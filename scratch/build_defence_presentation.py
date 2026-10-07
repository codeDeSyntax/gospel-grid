import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    template_path = r"j:\electron\wingridFYP\wingrid\public\FILES\UG SLIDE TEMPLATE.pptx"
    output_path = r"j:\electron\wingridFYP\wingrid\public\FILES\Wingrid_Final_Year_Defence_Presentation.pptx"

    prs = Presentation(template_path)
    
    # Remove existing blank placeholder slides except layout references
    # Note: in python-pptx we can clear slide contents or recreate slides using layouts
    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]
    
    # We will clear slides 1 to 8 and recreate a clean, complete slide deck
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]

    def add_title_slide(main_title, project_title, metadata_text):
        slide = prs.slides.add_slide(title_layout)
        for shape in slide.shapes:
            if shape.name.startswith("Title") or shape.placeholder_format.idx == 0:
                shape.text_frame.text = main_title
                for p in shape.text_frame.paragraphs:
                    p.font.name = "Arial"
                    p.font.size = Pt(22)
                    p.font.bold = True
                    p.font.color.rgb = RGBColor(0, 51, 102) # UG Navy Blue
            elif shape.name.startswith("Subtitle") or shape.placeholder_format.idx == 1:
                tf = shape.text_frame
                tf.text = project_title + "\n\n" + metadata_text
                for idx, p in enumerate(tf.paragraphs):
                    p.font.name = "Arial"
                    if idx == 0:
                        p.font.size = Pt(20)
                        p.font.bold = True
                        p.font.color.rgb = RGBColor(0, 96, 137) # Ocean Blue
                    else:
                        p.font.size = Pt(13)
                        p.font.color.rgb = RGBColor(50, 50, 50)
        return slide

    def add_content_slide(title_text, bullet_groups):
        slide = prs.slides.add_slide(content_layout)
        title_shape = None
        content_shape = None
        for shape in slide.shapes:
            if shape.has_text_frame and (shape.name.startswith("Title") or (shape.is_placeholder and shape.placeholder_format.idx == 0)):
                title_shape = shape
            elif shape.has_text_frame and (shape.name.startswith("Content") or (shape.is_placeholder and shape.placeholder_format.idx == 1)):
                content_shape = shape

        if title_shape:
            title_shape.text_frame.text = title_text
            for p in title_shape.text_frame.paragraphs:
                p.font.name = "Arial"
                p.font.size = Pt(22)
                p.font.bold = True
                p.font.color.rgb = RGBColor(0, 51, 102)

        if content_shape:
            tf = content_shape.text_frame
            tf.word_wrap = True
            tf.text = "" # clear default text
            
            first = True
            for heading, bullets in bullet_groups:
                if heading:
                    p = tf.paragraphs[0] if first else tf.add_paragraph()
                    first = False
                    p.text = heading
                    p.level = 0
                    p.font.name = "Arial"
                    p.font.size = Pt(15)
                    p.font.bold = True
                    p.font.color.rgb = RGBColor(0, 80, 120)
                    p.space_after = Pt(4)
                
                for b in bullets:
                    p = tf.paragraphs[0] if first else tf.add_paragraph()
                    first = False
                    p.text = b
                    p.level = 1 if heading else 0
                    p.font.name = "Arial"
                    p.font.size = Pt(13)
                    p.font.color.rgb = RGBColor(40, 40, 40)
                    p.space_after = Pt(4)
        return slide

    # SLIDE 1: Title Slide
    add_title_slide(
        "FINAL YEAR PROJECT DEFENCE & DEMO",
        "WINGRID: A Multi-Window Presentation Orchestration and Context-Aware Projection System for Live Display Environments",
        "University of Ghana - Department of Computer Science\nLevel 400 Final Year Project (v2.2.46)\n\nPresented by: [Student Name] (ID: [Student ID])\nSupervisor: [Supervisor Name]\nDate: October 2026"
    )

    # SLIDE 2: Presentation Outline
    add_content_slide(
        "Outline of Presentation",
        [
            ("Structure of the Presentation", [
                "1. Introduction & Research Motivation",
                "2. Problem Statement (The 4 Core Display Bottlenecks)",
                "3. Aims and Specific Objectives",
                "4. Related Work & Competitive Advantage Matrix",
                "5. Proposed Solution: 'Two Windows, One Operator' Architecture",
                "6. Key Subsystems & Real-Time AI Context Intelligence",
                "7. System Implementation & FYP Technical Scope",
                "8. Live Demonstration Workflow",
                "9. Conclusion & Future Research Directions"
            ])
        ]
    )

    # SLIDE 3: Introduction & Motivation
    add_content_slide(
        "Introduction & Background",
        [
            ("The Rise of Multi-Source Live Presentations", [
                "Modern presentation environments (churches, classrooms, auditoriums, hybrid streams) demand simultaneous live content from multiple software applications.",
                "Common live applications: presentation slides, Bible software, live code editors, web browsers, timers, and remote guest feeds.",
                "Presenters need to orchestrate and compose these sources dynamically in real time."
            ]),
            ("The Existing Technology Gap", [
                "Video conferencing tools (Zoom, Teams) were built for single-source remote meetings, not multi-display projection.",
                "Broadcast suites (OBS, vMix) are powerful but over-engineered for live operators with steep learning curves and heavy CPU overhead.",
                "There is no lightweight, zero-configuration system tailored for selective multi-window presentation orchestration."
            ])
        ]
    )

    # SLIDE 4: Problem Statement
    add_content_slide(
        "Problem Statement (The 4 Core Bottlenecks)",
        [
            ("1. Single-Source Sharing Restrictions", [
                "Platforms restrict sharing to one window or force sharing the entire screen, making side-by-side presentation impossible without manual window resizing."
            ]),
            ("2. Display Mode Conflicts & Screen Flickering", [
                "Different software requires different Windows display modes (Extended vs. Duplicate). Constantly toggling Win+P during live events causes flickering, resolution drops, and delays."
            ]),
            ("3. Privacy & Desktop Exposure", [
                "Sharing the full desktop leaks private notifications (WhatsApp, emails), desktop files, taskbars, and uncurated background windows to the audience."
            ]),
            ("4. Manual Content Augmentation Overhead", [
                "Displaying supporting context (scriptures, speaker lower-thirds, timers, quotes) requires dedicated crew manually typing or swapping OBS scenes."
            ])
        ]
    )

    # SLIDE 5: Aims and Objectives
    add_content_slide(
        "Aims and Objectives",
        [
            ("Primary Aim", [
                "To design, implement, and evaluate an intelligent desktop presentation orchestration system that makes live presentations more professional and smarter through selective multi-window composition and real-time AI intelligence."
            ]),
            ("Specific Objectives", [
                "1. Build a high-performance GPU-accelerated window capture pipeline with request coalescing (30-60 FPS).",
                "2. Develop a 'Two Windows, One Operator' dual-window display separation architecture.",
                "3. Implement an automated real-time AI speech-to-context producer using AssemblyAI and Groq Llama 3.1 (~200ms extraction).",
                "4. Engineer local mDNS device discovery and permission-gated WebRTC P2P remote screen streaming.",
                "5. Integrate enterprise-grade OS-level DPAPI cryptography for secure credential storage."
            ])
        ]
    )

    # SLIDE 6: Competitive Analysis
    add_content_slide(
        "How Wingrid Beats Existing Solutions",
        [
            ("Why Existing Alternatives Fall Short:", [
                "Zoom / Teams: Restricts to 1 window, risks full desktop leakage, lacks secondary display routing, cloud-dependent.",
                "OBS Studio / vMix: Complex scene graph configuration, high CPU/GPU overhead, no automated real-time speech AI cards.",
                "ProPresenter / EasyWorship: Expensive ($400-$1000+), slide-centric architecture is clunky for live external application windows.",
                "Windows Snap / FancyZones: Single-monitor canvas keeps titlebars and taskbars visible; forces Win+P display toggling."
            ]),
            ("Wingrid's Competitive Edge:", [
                "1-Click responsive multi-window grids (Single, Dual, Triple, Quad) with per-window Hide/Show toggles.",
                "Strict zero-desktop isolation: captures window textures directly via GPU compositor.",
                "Real-time AI Context Intelligence: auto-generates broadcast cards in ~200ms from live spoken words.",
                "100% local GPU execution with zero internet required for physical display projection."
            ])
        ]
    )

    # SLIDE 7: Proposed Solution & Architecture
    add_content_slide(
        "Proposed Solution: System Architecture",
        [
            ("'Two Windows, One Operator' Display Model", [
                "Primary Monitor (Control Studio): Operator manages window picker, confidence monitors, AI carousels, timers, and layout controls.",
                "Secondary Monitor (PublishedLayout): Pristine, frameless fullscreen canvas rendered directly on the target projector/LED display.",
                "State synchronization flows across low-latency Electron IPC channels (update-projection-state)."
            ]),
            ("Core Technology Stack", [
                "Frontend & UI: Electron 33, React 18, TypeScript, Tailwind CSS, Redux Toolkit, Framer Motion.",
                "Media & Video: Windows Graphics Capture (WGC), Web Audio/Video APIs (getUserMedia), Canvas GPU rendering.",
                "AI & Cloud: 16kHz PCM Audio Worklet, AssemblyAI WebSocket Streaming, Groq LPU Inference (Llama 3.1).",
                "Networking & Security: mDNS Discovery, WebSocket Signaling Server, WebRTC P2P Data/Media, Windows DPAPI."
            ])
        ]
    )

    # SLIDE 8: Key Subsystems (Window Engine & Display Routing)
    add_content_slide(
        "Key Subsystems: Window Engine & Display Routing",
        [
            ("1. Native Windows Graphics Capture (WGC) Enumerator", [
                "Discovers capturable top-level desktop application windows with live icons, metadata, and preview thumbnails.",
                "Includes search filtering, window pinning, and undo/redo selection history."
            ]),
            ("2. Dynamic AutoFit Multi-Grid Compositor", [
                "Automatically arranges selected windows into Single, Dual (side-by-side), Triple, and Quad (2x2) responsive grid layouts.",
                "Features instant per-window Hide/Show toggles without disturbing other projected sources."
            ]),
            ("3. GPU Video Pipeline & Request Coalescing", [
                "Eliminates DWM composition stalls through in-flight request coalescing registry for desktopCapturer.",
                "Keyboard live broadcast controls: F5 (Toggle Live Projection), F6 (Instant Blackout), F7 (Freeze Frame)."
            ])
        ]
    )

    # SLIDE 9: Key Subsystems (AI Context & Network Collaboration)
    add_content_slide(
        "Key Subsystems: AI Context & LAN Collaboration",
        [
            ("1. Real-Time AI Context Intelligence Producer (v2.2.46)", [
                "Continuous speech ingestion via 16kHz PCM audio stream to AssemblyAI WebSocket.",
                "Fast structured JSON extraction via Groq LPU (Llama 3.1) in ~200ms: auto-detects scriptures, speaker lower-thirds, quotes, and statistics.",
                "8+ responsive broadcast card variants with 1-click 'Push to Live' and 'Hide' controls."
            ]),
            ("2. Built-In Presentation Feature Tiles", [
                "Countdown Timers, Stopwatches, Clock tiles, Image Gallery projection, and Live Captions.",
                "Floating speech-to-text draggable orb with multi-display broadcast targeting."
            ]),
            ("3. Network-Aware LAN Discovery & WebRTC Collaboration", [
                "Zero-config mDNS device discovery across local Wi-Fi/Ethernet.",
                "Permission-gated request flow with direct peer-to-peer WebRTC streaming (zero cloud bandwidth)."
            ])
        ]
    )

    # SLIDE 10: Live Demonstration & Evaluation
    add_content_slide(
        "System Demonstration & Verification Workflow",
        [
            ("Demonstration Flow for the Panel:", [
                "Step 1: Launch Wingrid -> Automatically enumerate running desktop windows with live thumbnails.",
                "Step 2: Assign multiple windows (e.g., Slide Deck + Bible Software + Timer) into a Dual/Quad grid.",
                "Step 3: Trigger F5 Projection -> Verify pristine fullscreen output on secondary display with zero desktop clutter.",
                "Step 4: Speak into microphone -> Demonstrate AI Context Engine auto-generating scripture card in ~200ms -> 1-Click Push to Live.",
                "Step 5: Test live F6 (Blackout) and F7 (Freeze Frame) controls.",
                "Step 6: Initiate LAN Discovery -> Stream a remote screen feed via permission-gated WebRTC."
            ]),
            ("Performance & Stability Metrics:", [
                "Consistent 30-60 FPS GPU rendering with minimal CPU utilization (< 8%).",
                "Instant multi-display layout synchronization (< 16ms IPC latency)."
            ])
        ]
    )

    # SLIDE 11: Conclusion & Future Scope
    add_content_slide(
        "Conclusion & Future Scope",
        [
            ("Summary of Contributions", [
                "1. Successfully solved the 4 core display bottlenecks (single-source, display mode conflicts, privacy, manual overhead).",
                "2. Designed and implemented the 'Two Windows, One Operator' presentation model with GPU acceleration.",
                "3. Introduced real-time speech AI context intelligence and broadcast-grade presentation tiles."
            ]),
            ("Future Research & Enhancements", [
                "Multi-speaker voice diarization for panel discussions and multi-speaker conferences.",
                "Cross-platform display server support for macOS (ScreenCaptureKit) and Linux (Wayland/X11).",
                "Automated smart-focus camera tracking integrated with AI speaker lower-thirds."
            ])
        ]
    )

    # SLIDE 12: End Slide (Q&A)
    add_title_slide(
        "FINAL YEAR PROJECT DEFENCE & DEMO",
        "THANK YOU FOR YOUR ATTENTION",
        "Wingrid: A Multi-Window Presentation Orchestration and Context-Aware Projection System for Live Display Environments\n\nQuestions & Comments are Welcome\n\nDepartment of Computer Science, University of Ghana"
    )

    prs.save(output_path)
    print(f"Successfully generated presentation at: {output_path}")

if __name__ == "__main__":
    build_presentation()
